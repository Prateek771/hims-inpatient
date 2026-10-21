from pathlib import Path
import json, re
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

out = Path(__file__).resolve().parent
raw = (out / 'guide-draft.md').read_text(encoding='utf-8-sig')
raw = raw.replace('The nurse prepares Meera and gives her a physical checklist.', 'The nurse confirms readiness, handover and the bed-vacancy time used for room-charge cutoff, and gives Meera a physical checklist.')
raw = raw.replace('**Essential Admin sections:** Patients and Admissions, Beds, Payments and Billing, Discharge. Admission, Nurse, Pharmacy, Insurance and Admin/Billing handle their own work; the Doctor remains the clinical decision-maker.', '**Essential Admin sections:** Patients & admissions, Beds, Payments & insurance, Medicine bills & returns, Discharge, Staff & settings.')
for original, display in {'Rs120000':'Rs 1,20,000','Rs115000':'Rs 1,15,000','Rs90000':'Rs 90,000','Rs30000':'Rs 30,000','Rs25000':'Rs 25,000','Rs5000':'Rs 5,000','Rs0':'Rs 0'}.items():
    raw = raw.replace(original, display)
(out / 'guide-draft.md').write_text(raw, encoding='utf-8')
paras = [p.strip() for p in raw.split('\n\n') if p.strip()]
content = {
    'title': paras[0].removeprefix('# '),
    'intro': paras[1],
    'steps': paras[2:10],
    'money': paras[10],
    'admin': paras[11],
    'source': 'Sources: supplied Hospital Management patient journey and conversation. Later review is supplemental.'
}
(out / 'content.json').write_text(json.dumps(content, indent=2, ensure_ascii=False), encoding='utf-8')

# Editable Word text and the PDF use the same paragraphs.
def add_markdown(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    for i, piece in enumerate(text.split('**')):
        run = p.add_run(piece)
        run.bold = (i % 2 == 1)
    return p

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.52); sec.bottom_margin = Inches(.52)
sec.left_margin = Inches(.6); sec.right_margin = Inches(.6)
normal = doc.styles['Normal']
normal.font.name = 'Arial'; normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.05
normal.paragraph_format.space_after = Pt(7)
title = doc.styles['Title']
title.font.name = 'Arial'; title.font.size = Pt(19)
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(9)
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.keep_with_next = True
for border in list(title.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
p = doc.add_paragraph(content['title'], style='Title')
add_markdown(doc, content['intro'])
for text in content['steps']:
    p = add_markdown(doc, text)
    p.paragraph_format.keep_together = True
add_markdown(doc, content['money'])
add_markdown(doc, content['admin'])
p = doc.add_paragraph(content['source'])
p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(0)
for run in p.runs: run.font.size = Pt(8.5)
doc.core_properties.title = content['title']
doc.core_properties.subject = 'Plain language inpatient journey'
doc.core_properties.author = 'HIMS'
doc.save(out / 'hims-one-patient-admission-to-exit.docx')

pdfmetrics.registerFont(TTFont('Arial', r'C:\Windows\Fonts\arial.ttf'))
pdfmetrics.registerFont(TTFont('ArialBold', r'C:\Windows\Fonts\arialbd.ttf'))
pdfmetrics.registerFontFamily('Arial', normal='Arial', bold='ArialBold', italic='Arial', boldItalic='ArialBold')
body = ParagraphStyle('Body', fontName='Arial', fontSize=11, leading=13.3, spaceAfter=7, alignment=TA_LEFT)
title_style = ParagraphStyle('Title', parent=body, fontName='ArialBold', fontSize=19, leading=23, spaceAfter=10)
source_style = ParagraphStyle('Source', parent=body, fontSize=8.5, leading=10.5, spaceBefore=2, spaceAfter=0)
def pdf_text(text):
    pieces = text.split('**')
    return ''.join(('<b>'+escape(piece)+'</b>') if i % 2 else escape(piece) for i, piece in enumerate(pieces))
story = [Paragraph(escape(content['title']), title_style), Paragraph(pdf_text(content['intro']), body)]
for text in content['steps']: story.append(Paragraph(pdf_text(text), body))
story.extend([Paragraph(pdf_text(content['money']), body), Paragraph(pdf_text(content['admin']), body), Paragraph(escape(content['source']), source_style)])
SimpleDocTemplate(str(out / 'hims-one-patient-admission-to-exit.pdf'), pagesize=letter, rightMargin=43.2, leftMargin=43.2, topMargin=37.44, bottomMargin=37.44, title=content['title'], author='HIMS').build(story)
print('Word and PDF created from common content; words:', len(re.findall(r'\S+', ' '.join([content['title'], content['intro'], *content['steps'], content['money'], content['admin'], content['source']]))))
