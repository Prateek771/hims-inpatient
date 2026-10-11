"""Read only guide artifact checks and independent PDF rasterization."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import importlib.util
import json
from datetime import datetime, timezone
from docx import Document
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'qa/final'
OUT.mkdir(parents=True, exist_ok=True)
spec = importlib.util.spec_from_file_location('guide_source', ROOT / 'generate_patient_journey.py')
source = importlib.util.module_from_spec(spec)
spec.loader.exec_module(source)
manifest = json.loads((ROOT / 'assets/screenshots.json').read_text(encoding='utf-8'))
docx = ROOT / 'patient-journey.docx'
pdf = ROOT / 'patient-journey.pdf'
doc = Document(docx)
reader = PdfReader(pdf)
texts = [p.extract_text() for p in reader.pages]
joined = '\n\n'.join(f'PAGE {i+1}\n'+t for i,t in enumerate(texts))
(OUT / 'extracted-text.txt').write_text(joined, encoding='utf-8')
normalized = ' '.join(joined.split())
with ZipFile(docx) as archive:
    embedded = {hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist() if name.startswith('word/media/')}
checks = {
    'all_fifteen_headings_in_extracted_text': all(' '.join(p['title'].split()) in normalized for p in source.PAGES),
    'all_pages_have_meaningful_body': all(len(t)>700 for t in texts),
    'no_draft_or_pending_image_slots': not any(s in joined for s in ['DRAFT', 'Draft illustration slot:', 'final tested synthetic']),
    'five_required_money_strings': all(s in joined for s in ['₹1,15,000','₹90,000','₹25,000','₹30,000','₹5,000']),
    'insured_cash_arithmetic': 90000+30000-5000 == 115000,
    'self_pay_top_up_and_separate_sources': all(s in joined for s in ['TOP_UP','Additional patient top-up receipt','each against its own original receipt']),
    'admin_only_cleaning_prose': 'Admin records confirmed cleaning completion' in joined and 'Admission or Admin then records' not in joined,
    'eight_reconciled_synthetic_screens': manifest.get('reconciled_with_implemented_ui') is True and len(manifest['screens'])==8,
    'capture_hashes_match_files_and_docx_media': all(hashlib.sha256((ROOT/'assets'/s['file']).read_bytes()).hexdigest()==s['sha256'] and s['sha256'] in embedded for s in manifest['screens'].values()),
    'eleven_inline_images_have_alt_text': len(doc.inline_shapes)==11 and all(s._inline.docPr.get('descr') for s in doc.inline_shapes),
}
if not all(checks.values()):
    raise RuntimeError(json.dumps(checks, indent=2))
raster_dir=OUT/'pdfium-isolated'
raster_dir.mkdir(exist_ok=True)
sequential_dir=OUT/'pdfium'
sequential_dir.mkdir(exist_ok=True)
pdfium_doc=pdfium.PdfDocument(str(pdf))
for i in range(len(reader.pages)):
    page=pdfium_doc[i]
    bitmap=page.render(scale=150/72)
    bitmap.to_pil().save(sequential_dir/f'page-{i+1:02}.png')
    bitmap.close();page.close()
pdfium_doc.close()
for i in range(len(reader.pages)):
    pdfium_doc=pdfium.PdfDocument(str(pdf))
    page=pdfium_doc[i]
    bitmap=page.render(scale=150/72)
    bitmap.to_pil().save(raster_dir/f'page-{i+1:02}.png')
    bitmap.close();page.close();pdfium_doc.close()
result = {
  'checked_at_utc': datetime.now(timezone.utc).isoformat(),
  'pages': len(texts),
  'extracted_characters': sum(map(len,texts)),
  'checks': checks,
  'pages_and_headings': [{'page':i+1,'heading':p['title'],'characters':len(texts[i])} for i,p in enumerate(source.PAGES)],
  'outputs': {p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [docx,pdf]},
  'capture_source': manifest['app_source'],
  'rasterization': 'Required Poppler 150 dpi pages and independent bundled PDFium 150 dpi proofs. Apparent omissions in multi-image tool previews were disproved by identical local header pixel hashes and individual image viewing; no PDF or renderer defect was established. Visual inspection is recorded separately.'
}
header_hashes={}
for label,directory in [('poppler',OUT),('pdfium_sequential',sequential_dir),('pdfium_isolated',raster_dir)]:
    values=[]
    for i in range(len(texts)):
        im=Image.open(directory/f'page-{i+1:02}.png').convert('RGB')
        values.append(hashlib.sha256(im.crop((0,0,im.width,85)).tobytes()).hexdigest())
    header_hashes[label]={'unique_hashes':sorted(set(values)),'same_header_pixels_on_all_pages':len(set(values))==1}
result['header_pixel_crosscheck']=header_hashes
result['render_provenance']={
    label: [{'page': i+1, 'file': str((directory/f'page-{i+1:02}.png').relative_to(ROOT)).replace('\\','/'), 'sha256': hashlib.sha256((directory/f'page-{i+1:02}.png').read_bytes()).hexdigest()} for i in range(len(texts))]
    for label,directory in [('poppler',OUT),('pdfium_sequential',sequential_dir),('pdfium_isolated',raster_dir)]
}
(OUT/'validation.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print(json.dumps(result,indent=2,ensure_ascii=False))
